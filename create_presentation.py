import os
import sys
from PIL import Image, ImageDraw, ImageOps
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# --- Color Palette (Extracted from styles.css) ---
BG_COLOR = RGBColor(7, 16, 18)          # #071012 (Deep Dark Slate)
SURFACE_COLOR = RGBColor(14, 25, 28)     # #0E191C (Card background)
SURFACE_RAISED = RGBColor(19, 36, 40)   # #132428 (Elevated card)
BORDER_COLOR = RGBColor(32, 58, 62)     # #203A3E (Card borders)
CYAN = RGBColor(112, 240, 223)          # #70F0DF (Primary accent)
LIME = RGBColor(201, 247, 108)          # #C9F76C (Secondary accent)
TEXT_WHITE = RGBColor(237, 244, 237)    # #EDF4ED (Primary text)
TEXT_MUTED = RGBColor(148, 169, 164)    # #94A9A4 (Muted text)
TEXT_FAINT = RGBColor(94, 116, 111)     # #5E746F (Faint text)
BADGE_BG = RGBColor(16, 38, 40)         # Badge container background

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"
FONT_MONO = "Consolas"

def create_circular_profile_image(input_path, output_path, size=800):
    """Processes profile.jpg into a circular framed avatar with an elegant cyan ring."""
    try:
        im = Image.open(input_path).convert("RGBA")
        # Square crop
        min_dim = min(im.size)
        left = (im.width - min_dim) // 2
        top = (im.height - min_dim) // 2
        im = im.crop((left, top, left + min_dim, top + min_dim))
        im = im.resize((size, size), Image.Resampling.LANCZOS)

        # Create circular mask
        mask = Image.new("L", (size, size), 0)
        draw = ImageDraw.Draw(mask)
        draw.ellipse((0, 0, size, size), fill=255)

        # Apply mask
        output = Image.new("RGBA", (size, size), (0, 0, 0, 0))
        output.paste(im, (0, 0), mask=mask)

        # Draw double border rings
        border_draw = ImageDraw.Draw(output)
        border_draw.ellipse((4, 4, size - 5, size - 5), outline=(112, 240, 223, 180), width=6)
        border_draw.ellipse((14, 14, size - 15, size - 15), outline=(201, 247, 108, 90), width=3)

        output.save(output_path, format="PNG")
        return True
    except Exception as e:
        print(f"Error creating circular profile image: {e}")
        return False

def add_slide_background(slide, width=Inches(13.333), height=Inches(7.5)):
    """Draws a rich solid dark background on the slide."""
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, width, height
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = BG_COLOR
    bg_shape.line.fill.background() # No border
    return bg_shape

def add_header(slide, section_number, section_title, page_heading=None, page_subheading=None):
    """Adds a standard website-style top header with label and title."""
    # Top Section Label (e.g. "01 / ABOUT ME")
    label_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.45), Inches(11.5), Inches(0.4))
    tf = label_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"{section_number}  //  {section_title.upper()}"
    p.font.name = FONT_MONO
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Main Page Heading
    if page_heading:
        heading_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.85), Inches(11.5), Inches(0.8))
        tf2 = heading_box.text_frame
        tf2.word_wrap = True
        tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
        p2 = tf2.paragraphs[0]
        p2.text = page_heading
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE

        if page_subheading:
            p3 = tf2.add_paragraph()
            p3.text = page_subheading
            p3.font.name = FONT_BODY
            p3.font.size = Pt(13)
            p3.font.color.rgb = TEXT_MUTED
            p3.space_before = Pt(4)

def add_card(slide, left, top, width, height, bg_color=SURFACE_COLOR, border_color=BORDER_COLOR):
    """Adds a styled card container shape."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    return card

def add_footer_tag(slide):
    """Adds branding watermark at the bottom of the slide."""
    box = slide.shapes.add_textbox(Inches(0.9), Inches(6.9), Inches(11.5), Inches(0.35))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "OE.  Omar Ebied  |  Portfolio Presentation"
    p.font.name = FONT_MONO
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_FAINT

def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    base_dir = os.path.dirname(os.path.abspath(__file__))
    img_input = os.path.join(base_dir, "profile.jpg")
    img_processed = os.path.join(base_dir, "profile_circle.png")

    has_profile_img = False
    if os.path.exists(img_input):
        has_profile_img = create_circular_profile_image(img_input, img_processed)

    # =========================================================================
    # SLIDE 1: Title / Hero Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide1)

    # Header brand OE.
    brand_box = slide1.shapes.add_textbox(Inches(0.9), Inches(0.55), Inches(4), Inches(0.4))
    tf = brand_box.text_frame
    p = tf.paragraphs[0]
    p.text = "OE.  OMAR EBIED"
    p.font.name = FONT_MONO
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = LIME

    # Eyebrow status pill
    status_card = add_card(slide1, Inches(0.9), Inches(1.25), Inches(3.4), Inches(0.4), SURFACE_RAISED, CYAN)
    status_box = slide1.shapes.add_textbox(Inches(1.0), Inches(1.27), Inches(3.2), Inches(0.35))
    tf = status_box.text_frame
    p = tf.paragraphs[0]
    p.text = "●  AVAILABLE FOR OPPORTUNITIES"
    p.font.name = FONT_MONO
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Main Greeting & Name
    hero_box = slide1.shapes.add_textbox(Inches(0.9), Inches(1.85), Inches(7.6), Inches(4.5))
    tf = hero_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_greet = tf.paragraphs[0]
    p_greet.text = "Hello, I’m"
    p_greet.font.name = FONT_BODY
    p_greet.font.size = Pt(18)
    p_greet.font.color.rgb = TEXT_MUTED

    p_name = tf.add_paragraph()
    p_name.text = "Omar Ebied."
    p_name.font.name = FONT_HEADING
    p_name.font.size = Pt(50)
    p_name.font.bold = True
    p_name.font.color.rgb = TEXT_WHITE
    p_name.space_after = Pt(8)

    p_role = tf.add_paragraph()
    p_role.text = "Computer Science & IT Student  |  Aspiring Big Data Engineer"
    p_role.font.name = FONT_HEADING
    p_role.font.size = Pt(15.5)
    p_role.font.bold = True
    p_role.font.color.rgb = CYAN
    p_role.space_after = Pt(14)

    p_usp_en = tf.add_paragraph()
    p_usp_en.text = "I help growing businesses unlock actionable insights from their data and secure their systems through data-driven analysis and strong security fundamentals."
    p_usp_en.font.name = FONT_BODY
    p_usp_en.font.size = Pt(13)
    p_usp_en.font.color.rgb = TEXT_MUTED
    p_usp_en.space_after = Pt(12)

    p_usp_ar = tf.add_paragraph()
    p_usp_ar.text = "« أساعد الشركات على استخراج رؤى دقيقة من بياناتهم وتأمين أنظمتهم من خلال تحليل البيانات المعتمد على أسس أمان قوية »"
    p_usp_ar.font.name = "Segoe UI"
    p_usp_ar.font.size = Pt(12)
    p_usp_ar.font.italic = True
    p_usp_ar.font.color.rgb = LIME

    # Meta Tags row
    meta_box = slide1.shapes.add_textbox(Inches(0.9), Inches(6.05), Inches(7.5), Inches(0.5))
    tf_meta = meta_box.text_frame
    p_meta = tf_meta.paragraphs[0]
    p_meta.text = "📍 Cairo, Egypt   •   🏛 Helwan National University (HNU)   •   🎓 Class of 2028"
    p_meta.font.name = FONT_MONO
    p_meta.font.size = Pt(10.5)
    p_meta.font.color.rgb = TEXT_MUTED

    # Right side: Visual photo frame & floating chips
    photo_left = Inches(8.9)
    photo_top = Inches(1.35)
    photo_size = Inches(3.6)

    # Background visual aura card
    add_card(slide1, photo_left - Inches(0.15), photo_top - Inches(0.15), photo_size + Inches(0.3), photo_size + Inches(0.3), SURFACE_COLOR, BORDER_COLOR)

    if has_profile_img and os.path.exists(img_processed):
        slide1.shapes.add_picture(img_processed, photo_left, photo_top, photo_size, photo_size)
    else:
        # Fallback monogram
        avatar_card = add_card(slide1, photo_left, photo_top, photo_size, photo_size, SURFACE_RAISED, CYAN)
        av_box = slide1.shapes.add_textbox(photo_left, photo_top + Inches(1.2), photo_size, Inches(1.2))
        p_av = av_box.text_frame.paragraphs[0]
        p_av.text = "OE"
        p_av.font.name = FONT_HEADING
        p_av.font.size = Pt(56)
        p_av.font.bold = True
        p_av.font.color.rgb = CYAN
        p_av.alignment = PP_ALIGN.CENTER

    # Floating Chips
    chip1 = add_card(slide1, Inches(8.3), Inches(1.9), Inches(1.6), Inches(0.45), SURFACE_RAISED, LIME)
    c1_box = slide1.shapes.add_textbox(Inches(8.3), Inches(1.92), Inches(1.6), Inches(0.4))
    p_c1 = c1_box.text_frame.paragraphs[0]
    p_c1.text = "⚡ DATA ENG"
    p_c1.font.name = FONT_MONO
    p_c1.font.size = Pt(9.5)
    p_c1.font.bold = True
    p_c1.font.color.rgb = LIME
    p_c1.alignment = PP_ALIGN.CENTER

    chip2 = add_card(slide1, Inches(11.0), Inches(4.5), Inches(1.6), Inches(0.45), SURFACE_RAISED, CYAN)
    c2_box = slide1.shapes.add_textbox(Inches(11.0), Inches(4.52), Inches(1.6), Inches(0.4))
    p_c2 = c2_box.text_frame.paragraphs[0]
    p_c2.text = "🛡 SECURE"
    p_c2.font.name = FONT_MONO
    p_c2.font.size = Pt(9.5)
    p_c2.font.bold = True
    p_c2.font.color.rgb = CYAN
    p_c2.alignment = PP_ALIGN.CENTER

    add_footer_tag(slide1)

    # =========================================================================
    # SLIDE 2: About Me & Background
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide2)
    add_header(slide2, "01", "About Me", "Building the bridge between raw data and clear decisions.", "Developing the technical foundation to turn complex information into reliable systems.")

    # Left Narrative Card
    left_w = Inches(6.8)
    add_card(slide2, Inches(0.9), Inches(2.0), left_w, Inches(4.6), SURFACE_COLOR, BORDER_COLOR)
    about_box = slide2.shapes.add_textbox(Inches(1.2), Inches(2.3), left_w - Inches(0.6), Inches(4.0))
    tf_about = about_box.text_frame
    tf_about.word_wrap = True
    tf_about.margin_left = tf_about.margin_top = tf_about.margin_right = tf_about.margin_bottom = 0

    p_a1 = tf_about.paragraphs[0]
    p_a1.text = "I’m a Computer Science & Information Technology student at Helwan National University, developing the technical foundation to turn complex information into useful, reliable systems."
    p_a1.font.name = FONT_BODY
    p_a1.font.size = Pt(14)
    p_a1.font.color.rgb = TEXT_WHITE
    p_a1.space_after = Pt(14)

    p_a2 = tf_about.add_paragraph()
    p_a2.text = "My core focus sits at the intersection of Data Engineering, Big Data, and Cybersecurity—where well-designed data foundations and strong security practices empower organizations to scale with confidence."
    p_a2.font.name = FONT_BODY
    p_a2.font.size = Pt(13.5)
    p_a2.font.color.rgb = TEXT_MUTED
    p_a2.space_after = Pt(16)

    p_a3 = tf_about.add_paragraph()
    p_a3.text = "Key Philosophies:"
    p_a3.font.name = FONT_MONO
    p_a3.font.size = Pt(11.5)
    p_a3.font.bold = True
    p_a3.font.color.rgb = CYAN
    p_a3.space_after = Pt(6)

    p_a4 = tf_about.add_paragraph()
    p_a4.text = "• Data Accuracy: Rigorous validation & scalable relational schema design.\n• Threat Resilience: Deep awareness of network risk, attacks, & defense.\n• Continuous Growth: Hands-on labs, competitive training programs, & open source."
    p_a4.font.name = FONT_BODY
    p_a4.font.size = Pt(12)
    p_a4.font.color.rgb = TEXT_MUTED

    # Right KPI Stats Cards (3 cards vertically)
    right_x = Inches(8.0)
    card_w = Inches(4.4)
    kpis = [
        ("2024 — 2028", "CSIT Degree at HNU", "4-year academic program specializing in Computer Science & Information Technology.", LIME),
        ("2+ Technical Tracks", "DEPI & NTI Trainee", "Intensive training in Microsoft Data Engineering and Network Cybersecurity.", CYAN),
        ("3 Core Projects", "Full-Stack & Security Labs", "Proven hands-on implementations in relational databases, network attacks, and web systems.", TEXT_WHITE)
    ]

    for i, (val, title, desc, color) in enumerate(kpis):
        top_y = Inches(2.0) + Inches(i * 1.58)
        add_card(slide2, right_x, top_y, card_w, Inches(1.42), SURFACE_COLOR, BORDER_COLOR)
        kpi_box = slide2.shapes.add_textbox(right_x + Inches(0.3), top_y + Inches(0.18), card_w - Inches(0.6), Inches(1.1))
        tf_k = kpi_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0

        p1 = tf_k.paragraphs[0]
        p1.text = val
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(20)
        p1.font.bold = True
        p1.font.color.rgb = color

        p2 = tf_k.add_paragraph()
        p2.text = title
        p2.font.name = FONT_MONO
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(2)

        p3 = tf_k.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(10)
        p3.font.color.rgb = TEXT_MUTED
        p3.space_before = Pt(2)

    add_footer_tag(slide2)

    # =========================================================================
    # SLIDE 3: Education & Academic Foundation
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide3)
    add_header(slide3, "02", "Education", "Academic Foundation & Core Coursework", "Structured university degree paired with practical hands-on engineering principles.")

    # Main University Card (Hero Card)
    main_edu = add_card(slide3, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.2), SURFACE_COLOR, CYAN)
    stamp_card = add_card(slide3, Inches(1.25), Inches(2.35), Inches(1.5), Inches(1.5), SURFACE_RAISED, LIME)
    s_box = slide3.shapes.add_textbox(Inches(1.25), Inches(2.75), Inches(1.5), Inches(0.7))
    p_s = s_box.text_frame.paragraphs[0]
    p_s.text = "HNU"
    p_s.font.name = FONT_HEADING
    p_s.font.size = Pt(26)
    p_s.font.bold = True
    p_s.font.color.rgb = LIME
    p_s.alignment = PP_ALIGN.CENTER

    edu_box = slide3.shapes.add_textbox(Inches(3.1), Inches(2.25), Inches(9.0), Inches(1.7))
    tf_edu = edu_box.text_frame
    tf_edu.word_wrap = True
    tf_edu.margin_left = tf_edu.margin_top = tf_edu.margin_right = tf_edu.margin_bottom = 0

    p_kicker = tf_edu.paragraphs[0]
    p_kicker.text = "2024 — 2028  •  CAIRO, EGYPT  •  UNDERGRADUATE"
    p_kicker.font.name = FONT_MONO
    p_kicker.font.size = Pt(11)
    p_kicker.font.bold = True
    p_kicker.font.color.rgb = CYAN

    p_deg = tf_edu.add_paragraph()
    p_deg.text = "Bachelor of Computer Science & Information Technology (CSIT)"
    p_deg.font.name = FONT_HEADING
    p_deg.font.size = Pt(21)
    p_deg.font.bold = True
    p_deg.font.color.rgb = TEXT_WHITE
    p_deg.space_before = Pt(4)

    p_uni = tf_edu.add_paragraph()
    p_uni.text = "Helwan National University (HNU)"
    p_uni.font.name = FONT_BODY
    p_uni.font.size = Pt(14)
    p_uni.font.color.rgb = TEXT_MUTED
    p_uni.space_before = Pt(3)

    # Coursework Badges / Grid (5 Subjects)
    cw_heading = slide3.shapes.add_textbox(Inches(0.9), Inches(4.45), Inches(11.5), Inches(0.4))
    p_cw = cw_heading.text_frame.paragraphs[0]
    p_cw.text = "SELECTED ACADEMIC COURSEWORK"
    p_cw.font.name = FONT_MONO
    p_cw.font.size = Pt(11)
    p_cw.font.bold = True
    p_cw.font.color.rgb = LIME

    courses = [
        ("Database Systems", "Relational modeling, SQL queries, normalization, ACID properties, transaction processing."),
        ("Algorithms & DS", "Complexity analysis, sorting, graph traversal, search algorithms, data structures."),
        ("Computer Networks", "OSI model, TCP/IP, IP routing, subnetting, Wireshark packet analysis, protocols."),
        ("Operating Systems", "Process scheduling, memory management, file systems, concurrency, Linux CLI."),
        ("Programming", "Object-oriented programming, data encapsulation, C/C++, Java, and Python fundamentals.")
    ]

    card_width = Inches(2.18)
    for i, (title, desc) in enumerate(courses):
        card_x = Inches(0.9) + Inches(i * 2.33)
        add_card(slide3, card_x, Inches(4.9), card_width, Inches(1.85), SURFACE_COLOR, BORDER_COLOR)
        c_box = slide3.shapes.add_textbox(card_x + Inches(0.18), Inches(5.05), card_width - Inches(0.36), Inches(1.55))
        tf_c = c_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0

        p1 = tf_c.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = CYAN

        p2 = tf_c.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_BODY
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED
        p2.space_before = Pt(6)

    add_footer_tag(slide3)

    # =========================================================================
    # SLIDE 4: Technical Skills & Toolkit
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide4)
    add_header(slide4, "03", "Skills", "A practical, growing technical toolkit.", "Hands-on capabilities developed through academic study, specialized tracks, and project building.")

    skills_data = [
        ("01", "Programming", "Python · C · C++ · Java", [
            "Object-Oriented Programming (OOP)",
            "Data structures & algorithms",
            "Scripting & automation with Python",
            "System programming fundamentals"
        ], CYAN),
        ("02", "Data & Databases", "SQL · MS SQL Server · RDBMS", [
            "Database schema & relational modeling",
            "Complex SQL querying & optimization",
            "Data pipelines & ETL concepts",
            "Big Data fundamentals"
        ], LIME),
        ("03", "Cybersecurity & Net", "CCNA · Kali · Wireshark", [
            "Network protocols & packet capture",
            "Traffic analysis & vulnerability scanning",
            "Man-in-the-Middle (MITM) simulations",
            "Security assessment & reconnaissance"
        ], CYAN),
        ("04", "Web Technologies", "HTML · CSS · JavaScript", [
            "Responsive & accessible UI layout",
            "Modern CSS Grid, Flexbox, variables",
            "Interactive DOM scripting",
            "Client-server communication"
        ], TEXT_WHITE),
        ("05", "Soft Skills", "Problem-Solving · Teamwork", [
            "Analytical thinking & debugging",
            "Collaborative teamwork & git flow",
            "Adaptability in fast-evolving tech",
            "Technical documentation & presentation"
        ], LIME)
    ]

    col_w = Inches(2.18)
    for i, (num, cat_name, tech_line, items, accent) in enumerate(skills_data):
        col_x = Inches(0.9) + Inches(i * 2.33)
        add_card(slide4, col_x, Inches(2.0), col_w, Inches(4.6), SURFACE_COLOR, BORDER_COLOR)

        s_box = slide4.shapes.add_textbox(col_x + Inches(0.2), Inches(2.2), col_w - Inches(0.4), Inches(4.2))
        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0

        p_num = tf_s.paragraphs[0]
        p_num.text = num
        p_num.font.name = FONT_MONO
        p_num.font.size = Pt(13)
        p_num.font.bold = True
        p_num.font.color.rgb = accent

        p_name = tf_s.add_paragraph()
        p_name.text = cat_name
        p_name.font.name = FONT_HEADING
        p_name.font.size = Pt(14)
        p_name.font.bold = True
        p_name.font.color.rgb = TEXT_WHITE
        p_name.space_before = Pt(4)

        p_tech = tf_s.add_paragraph()
        p_tech.text = tech_line
        p_tech.font.name = FONT_MONO
        p_tech.font.size = Pt(9.5)
        p_tech.font.color.rgb = CYAN
        p_tech.space_before = Pt(4)
        p_tech.space_after = Pt(14)

        for item in items:
            p_item = tf_s.add_paragraph()
            p_item.text = f"• {item}"
            p_item.font.name = FONT_BODY
            p_item.font.size = Pt(10.5)
            p_item.font.color.rgb = TEXT_MUTED
            p_item.space_before = Pt(5)

    add_footer_tag(slide4)

    # =========================================================================
    # SLIDE 5: Experience & Professional Training
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide5)
    add_header(slide5, "04", "Experience", "Learning by solving real problems.", "Structured around the CAR (Challenge — Action — Result) framework.")

    experiences = [
        {
            "badge": "JUL 2026 — PRESENT  •  DATA SCIENCE",
            "title": "Microsoft Data Engineer Track",
            "org": "DEPI Trainee · Digital Egypt Pioneers Initiative",
            "c": "Building the technical foundations needed to engineer and operate modern data pipelines and enterprise analytics.",
            "a": "Training in core data engineering concepts, database architectures, analytics workflows, and Microsoft-aligned tools.",
            "r": "Strengthening a structured, industry-relevant competency path toward enterprise Big Data engineering.",
            "color": CYAN
        },
        {
            "badge": "SEP 2025  •  CYBERSECURITY",
            "title": "Cybersecurity & Network Defense",
            "org": "NTI Trainee & Freelancer · National Telecommunication Institute",
            "c": "Exploring practical network vulnerabilities, common attack vectors, and hands-on defense mechanisms.",
            "a": "Applied reconnaissance, traffic sniffing, security assessment, and attack simulations in isolated lab environments.",
            "r": "Built practical awareness of security fundamentals and how defensive controls protect mission-critical data.",
            "color": LIME
        }
    ]

    card_w = Inches(5.6)
    for i, exp in enumerate(experiences):
        card_x = Inches(0.9) + Inches(i * 5.9)
        add_card(slide5, card_x, Inches(2.0), card_w, Inches(4.6), SURFACE_COLOR, BORDER_COLOR)

        e_box = slide5.shapes.add_textbox(card_x + Inches(0.35), Inches(2.3), card_w - Inches(0.7), Inches(4.0))
        tf_e = e_box.text_frame
        tf_e.word_wrap = True
        tf_e.margin_left = tf_e.margin_top = tf_e.margin_right = tf_e.margin_bottom = 0

        p_b = tf_e.paragraphs[0]
        p_b.text = exp["badge"]
        p_b.font.name = FONT_MONO
        p_b.font.size = Pt(10.5)
        p_b.font.bold = True
        p_b.font.color.rgb = exp["color"]

        p_t = tf_e.add_paragraph()
        p_t.text = exp["title"]
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(20)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_before = Pt(4)

        p_o = tf_e.add_paragraph()
        p_o.text = exp["org"]
        p_o.font.name = FONT_BODY
        p_o.font.size = Pt(12)
        p_o.font.color.rgb = TEXT_MUTED
        p_o.space_after = Pt(14)

        # C - A - R blocks
        car_items = [
            ("C", "Challenge", exp["c"], CYAN),
            ("A", "Action", exp["a"], LIME),
            ("R", "Result", exp["r"], TEXT_WHITE)
        ]

        for code, label, txt, col in car_items:
            p_lbl = tf_e.add_paragraph()
            p_lbl.text = f"[{code}] {label}:"
            p_lbl.font.name = FONT_MONO
            p_lbl.font.size = Pt(11)
            p_lbl.font.bold = True
            p_lbl.font.color.rgb = col
            p_lbl.space_before = Pt(6)

            p_txt = tf_e.add_paragraph()
            p_txt.text = txt
            p_txt.font.name = FONT_BODY
            p_txt.font.size = Pt(11)
            p_txt.font.color.rgb = TEXT_MUTED
            p_txt.space_before = Pt(2)

    add_footer_tag(slide5)

    # =========================================================================
    # SLIDE 6: Services & Core Capabilities
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide6)
    add_header(slide6, "05", "Services", "Useful technical work, without the noise.", "Direct solutions addressing critical organizational data and security challenges.")

    services = [
        {
            "icon": "⌁",
            "name": "Data Analysis & Engineering",
            "problem": "Business data is scattered, disorganized, or difficult to extract actionable insights from.",
            "solution": "Designing structured pipelines that transform raw data into clear, reliable metrics and informed business decisions.",
            "outcomes": ["Clean ETL processing", "Interactive dashboards", "Actionable analytics"],
            "accent": CYAN
        },
        {
            "icon": "▣",
            "name": "Database Architecture & Design",
            "problem": "Information lacks a robust, organized home, causing redundancy and query bottlenecks.",
            "solution": "Engineering normalized relational database schemas (RDBMS) built for data integrity, high consistency, and scale.",
            "outcomes": ["Normalized schemas (3NF)", "Optimized SQL queries", "ACID transactional safety"],
            "accent": LIME
        },
        {
            "icon": "◎",
            "name": "Network & Security Assessment",
            "problem": "Basic network vulnerabilities and unmonitored ports remain unnoticed, creating risk.",
            "solution": "Delivering practical security fundamentals assessments to uncover risks and prioritize defense improvements.",
            "outcomes": ["Traffic reconnaissance", "Vulnerability auditing", "Security hardening recommendations"],
            "accent": CYAN
        }
    ]

    card_w = Inches(3.68)
    for i, srv in enumerate(services):
        card_x = Inches(0.9) + Inches(i * 3.9)
        add_card(slide6, card_x, Inches(2.0), card_w, Inches(4.6), SURFACE_COLOR, BORDER_COLOR)

        s_box = slide6.shapes.add_textbox(card_x + Inches(0.3), Inches(2.25), card_w - Inches(0.6), Inches(4.1))
        tf_s = s_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0

        p_ic = tf_s.paragraphs[0]
        p_ic.text = srv["icon"]
        p_ic.font.name = FONT_HEADING
        p_ic.font.size = Pt(28)
        p_ic.font.color.rgb = srv["accent"]

        p_t = tf_s.add_paragraph()
        p_t.text = srv["name"]
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(17)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_before = Pt(4)
        p_t.space_after = Pt(12)

        p_pr_lbl = tf_s.add_paragraph()
        p_pr_lbl.text = "THE PROBLEM:"
        p_pr_lbl.font.name = FONT_MONO
        p_pr_lbl.font.size = Pt(10)
        p_pr_lbl.font.bold = True
        p_pr_lbl.font.color.rgb = LIME

        p_pr = tf_s.add_paragraph()
        p_pr.text = srv["problem"]
        p_pr.font.name = FONT_BODY
        p_pr.font.size = Pt(11)
        p_pr.font.color.rgb = TEXT_MUTED
        p_pr.space_after = Pt(10)

        p_sol_lbl = tf_s.add_paragraph()
        p_sol_lbl.text = "THE RESULT:"
        p_sol_lbl.font.name = FONT_MONO
        p_sol_lbl.font.size = Pt(10)
        p_sol_lbl.font.bold = True
        p_sol_lbl.font.color.rgb = CYAN

        p_sol = tf_s.add_paragraph()
        p_sol.text = srv["solution"]
        p_sol.font.name = FONT_BODY
        p_sol.font.size = Pt(11)
        p_sol.font.color.rgb = TEXT_WHITE
        p_sol.space_after = Pt(10)

        for out in srv["outcomes"]:
            p_o = tf_s.add_paragraph()
            p_o.text = f"✔ {out}"
            p_o.font.name = FONT_BODY
            p_o.font.size = Pt(10)
            p_o.font.color.rgb = TEXT_MUTED

    add_footer_tag(slide6)

    # =========================================================================
    # SLIDE 7: Featured Projects (Selected Work)
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide7)
    add_header(slide7, "06", "Projects", "Selected work with purpose & impact.", "Practical implementations spanning databases, cybersecurity labs, and responsive web systems.")

    projects = [
        {
            "num": "PROJECT 01",
            "cat": "DATA & DATABASES",
            "title": "SQL Database Project",
            "tools": "MS SQL Server · Relational Modeling · T-SQL",
            "problem": "Organize connected business entities in a scalable, strictly normalized format.",
            "solution": "Engineered a relational database schema with structured primary/foreign keys, complex JOIN queries, and constraints.",
            "result": "Reliable data foundation ensuring zero duplication and fast querying.",
            "accent": CYAN
        },
        {
            "num": "PROJECT 02",
            "cat": "CYBERSECURITY LAB",
            "title": "MITM Lab (Man-in-the-Middle)",
            "tools": "Kali Linux · Wireshark · mitmproxy · ARP Spoofing",
            "problem": "Examine how unencrypted network traffic and spoofed ARP frames compromise integrity.",
            "solution": "Simulated ARP cache poisoning in a controlled sandbox to observe packet flow and credential leakage.",
            "result": "Actionable understanding of attack vectors and essential encryption safeguards.",
            "accent": LIME
        },
        {
            "num": "PROJECT 03",
            "cat": "WEB DEVELOPMENT",
            "title": "Online Book Store Website",
            "tools": "HTML5 · CSS3 · Modern JavaScript · Responsive UI",
            "problem": "Build an engaging, accessible catalog browsing experience for book lovers.",
            "solution": "Developed an interactive client-side web interface with catalog filtering and responsive design.",
            "result": "Smooth, polished front-end product demonstrating modern web engineering standards.",
            "accent": TEXT_WHITE
        }
    ]

    card_w = Inches(3.68)
    for i, proj in enumerate(projects):
        card_x = Inches(0.9) + Inches(i * 3.9)
        add_card(slide7, card_x, Inches(2.0), card_w, Inches(4.6), SURFACE_COLOR, BORDER_COLOR)

        p_box = slide7.shapes.add_textbox(card_x + Inches(0.3), Inches(2.2), card_w - Inches(0.6), Inches(4.2))
        tf_p = p_box.text_frame
        tf_p.word_wrap = True
        tf_p.margin_left = tf_p.margin_top = tf_p.margin_right = tf_p.margin_bottom = 0

        p_head = tf_p.paragraphs[0]
        p_head.text = f"{proj['num']}  •  {proj['cat']}"
        p_head.font.name = FONT_MONO
        p_head.font.size = Pt(10)
        p_head.font.bold = True
        p_head.font.color.rgb = proj["accent"]

        p_title = tf_p.add_paragraph()
        p_title.text = proj["title"]
        p_title.font.name = FONT_HEADING
        p_title.font.size = Pt(17)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.space_before = Pt(4)

        p_tools = tf_p.add_paragraph()
        p_tools.text = proj["tools"]
        p_tools.font.name = FONT_MONO
        p_tools.font.size = Pt(9.5)
        p_tools.font.color.rgb = CYAN
        p_tools.space_before = Pt(4)
        p_tools.space_after = Pt(10)

        # Problem / Solution / Result
        sections = [
            ("Problem", proj["problem"]),
            ("Solution", proj["solution"]),
            ("Result", proj["result"])
        ]
        for sec_name, sec_val in sections:
            p_sn = tf_p.add_paragraph()
            p_sn.text = f"{sec_name}:"
            p_sn.font.name = FONT_HEADING
            p_sn.font.size = Pt(10.5)
            p_sn.font.bold = True
            p_sn.font.color.rgb = LIME
            p_sn.space_before = Pt(4)

            p_sv = tf_p.add_paragraph()
            p_sv.text = sec_val
            p_sv.font.name = FONT_BODY
            p_sv.font.size = Pt(10.5)
            p_sv.font.color.rgb = TEXT_MUTED

    add_footer_tag(slide7)

    # =========================================================================
    # SLIDE 8: Achievements & Milestones
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide8)
    add_header(slide8, "07", "Achievements", "Milestones that keep me moving forward.", "Key accomplishments reflecting technical commitment, specialized training, and dedication.")

    achievements = [
        {
            "icon": "✦",
            "kicker": "DATA ENGINEERING",
            "title": "DEPI Selection",
            "sub": "Microsoft Data Engineer Track",
            "desc": "Competitively selected for the prestigious Digital Egypt Pioneers Initiative (DEPI). Undergoing intensive specialized training in big data pipelines, enterprise analytics, and cloud databases.",
            "stat": "Top Tier",
            "stat_label": "Selection Cohort",
            "color": CYAN
        },
        {
            "icon": "◈",
            "kicker": "CYBERSECURITY ACADEMY",
            "title": "NTI Cybersecurity Certification",
            "sub": "National Telecommunication Institute",
            "desc": "Successfully completed an intensive 60 technical hours cybersecurity training academy covering networking, reconnaissance, vulnerability assessment, and defense fundamentals.",
            "stat": "60 Hours",
            "stat_label": "Technical Training",
            "color": LIME
        },
        {
            "icon": "⌘",
            "kicker": "OPEN SOURCE & BUILDING",
            "title": "Build in Public & GitHub",
            "sub": "Independent Code Architecture",
            "desc": "Continuously applying computer science theory into real-world code. Publishing and documenting codebases, security labs, and database designs on GitHub for transparency and collaboration.",
            "stat": "3+ Labs",
            "stat_label": "Hands-On Repos",
            "color": TEXT_WHITE
        }
    ]

    card_w = Inches(3.68)
    for i, ach in enumerate(achievements):
        card_x = Inches(0.9) + Inches(i * 3.9)
        add_card(slide8, card_x, Inches(2.0), card_w, Inches(4.6), SURFACE_COLOR, BORDER_COLOR)

        a_box = slide8.shapes.add_textbox(card_x + Inches(0.3), Inches(2.25), card_w - Inches(0.6), Inches(4.1))
        tf_a = a_box.text_frame
        tf_a.word_wrap = True
        tf_a.margin_left = tf_a.margin_top = tf_a.margin_right = tf_a.margin_bottom = 0

        p_ic = tf_a.paragraphs[0]
        p_ic.text = f"{ach['icon']}  {ach['kicker']}"
        p_ic.font.name = FONT_MONO
        p_ic.font.size = Pt(10.5)
        p_ic.font.bold = True
        p_ic.font.color.rgb = ach["color"]

        p_t = tf_a.add_paragraph()
        p_t.text = ach["title"]
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(18)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_before = Pt(6)

        p_sub = tf_a.add_paragraph()
        p_sub.text = ach["sub"]
        p_sub.font.name = FONT_BODY
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = CYAN
        p_sub.space_before = Pt(2)
        p_sub.space_after = Pt(12)

        p_desc = tf_a.add_paragraph()
        p_desc.text = ach["desc"]
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = TEXT_MUTED
        p_desc.space_after = Pt(16)

        # Highlight Stat Pill
        p_st = tf_a.add_paragraph()
        p_st.text = f"★ {ach['stat']} — {ach['stat_label']}"
        p_st.font.name = FONT_MONO
        p_st.font.size = Pt(11)
        p_st.font.bold = True
        p_st.font.color.rgb = LIME

    add_footer_tag(slide8)

    # =========================================================================
    # SLIDE 9: Contact & Next Steps
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_slide_background(slide9)
    add_header(slide9, "08", "Contact", "Let’s make your data work harder.", "Available for internships, technical engineering opportunities, and project collaborations.")

    # Left Big Callout
    left_w = Inches(5.6)
    add_card(slide9, Inches(0.9), Inches(2.0), left_w, Inches(4.6), SURFACE_COLOR, CYAN)
    c_left_box = slide9.shapes.add_textbox(Inches(1.2), Inches(2.35), left_w - Inches(0.6), Inches(4.0))
    tf_cl = c_left_box.text_frame
    tf_cl.word_wrap = True
    tf_cl.margin_left = tf_cl.margin_top = tf_cl.margin_right = tf_cl.margin_bottom = 0

    p_c1 = tf_cl.paragraphs[0]
    p_c1.text = "Get in Touch"
    p_c1.font.name = FONT_HEADING
    p_c1.font.size = Pt(26)
    p_c1.font.bold = True
    p_c1.font.color.rgb = TEXT_WHITE

    p_c2 = tf_cl.add_paragraph()
    p_c2.text = "Have a project, internship, or technical collaboration in mind? I’d be glad to discuss how my data engineering and cybersecurity foundations can create value for your team."
    p_c2.font.name = FONT_BODY
    p_c2.font.size = Pt(13)
    p_c2.font.color.rgb = TEXT_MUTED
    p_c2.space_before = Pt(10)
    p_c2.space_after = Pt(18)

    p_quote = tf_cl.add_paragraph()
    p_quote.text = "“Data-minded. Security-aware. Always building.”"
    p_quote.font.name = FONT_HEADING
    p_quote.font.size = Pt(15)
    p_quote.font.bold = True
    p_quote.font.color.rgb = LIME
    p_quote.space_after = Pt(12)

    p_ar_quote = tf_cl.add_paragraph()
    p_ar_quote.text = "متاح للفرص والتدريبات ومشاريع هندسة البيانات وأمن المعلومات."
    p_ar_quote.font.name = "Segoe UI"
    p_ar_quote.font.size = Pt(12)
    p_ar_quote.font.color.rgb = CYAN

    # Right Contact Cards
    right_x = Inches(6.8)
    right_w = Inches(5.6)
    contacts = [
        ("EMAIL ADDRESS", "momr09989@gmail.com", "Direct correspondence for inquiries and project discussions", "✉", LIME),
        ("PHONE NUMBER", "01121024708  /  (+20) 1121024708", "Call or WhatsApp for immediate communication", "✆", CYAN),
        ("PROFESSIONAL PROFILES", "LinkedIn  •  GitHub", "Connect for career updates and inspect open source repositories", "◈", TEXT_WHITE),
        ("LOCATION & STATUS", "Cairo, Egypt  •  Helwan National University", "Open to on-site, hybrid, and remote opportunities", "⌖", LIME)
    ]

    for i, (kicker, val, note, icon, col) in enumerate(contacts):
        top_y = Inches(2.0) + Inches(i * 1.18)
        add_card(slide9, right_x, top_y, right_w, Inches(1.05), SURFACE_COLOR, BORDER_COLOR)

        ct_box = slide9.shapes.add_textbox(right_x + Inches(0.3), top_y + Inches(0.12), right_w - Inches(0.6), Inches(0.85))
        tf_ct = ct_box.text_frame
        tf_ct.word_wrap = True
        tf_ct.margin_left = tf_ct.margin_top = tf_ct.margin_right = tf_ct.margin_bottom = 0

        p1 = tf_ct.paragraphs[0]
        p1.text = f"{icon}  {kicker}"
        p1.font.name = FONT_MONO
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = col

        p2 = tf_ct.add_paragraph()
        p2.text = val
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(13.5)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(2)

        p3 = tf_ct.add_paragraph()
        p3.text = note
        p3.font.name = FONT_BODY
        p3.font.size = Pt(9.5)
        p3.font.color.rgb = TEXT_MUTED

    add_footer_tag(slide9)

    # Save output presentation
    output_pptx = os.path.join(base_dir, "Omar_Ebied_Portfolio.pptx")
    prs.save(output_pptx)
    print(f"Presentation successfully created at: {output_pptx}")

if __name__ == "__main__":
    main()
